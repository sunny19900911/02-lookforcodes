"""Read-only saved-XLSX validation using the bundled Python runtime."""
import collections
import hashlib
import json
import pathlib
import sys
import zipfile
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')

root = pathlib.Path(__file__).resolve().parent
book = root / 'CECS可下载PDF总表.xlsx'
master = json.loads((root / 'cecs-master.json').read_text(encoding='utf-8'))
with zipfile.ZipFile(book) as archive:
    assert archive.testzip() is None
workbook = openpyxl.load_workbook(book, data_only=False)
build = json.loads((root / 'workbook-build-receipt.json').read_text(encoding='utf-8'))
assert workbook.sheetnames == [s['name'] for s in build['sheets']]
expected = {'规范总表':len(master['standards']), '已验完整PDF':master['coverage']['verified_complete_pdf_links'], '链接明细':len(master['links']), '原始映射':81, '待查与问题':len(master['issues']), '状态与来源依据':len(master['evidence']), '变更记录':len(master['change_log'])}
if '完整草案PDF' in workbook.sheetnames:
    expected['完整草案PDF'] = master['coverage']['verified_complete_draft_pdf_links']
if '地区下载索引' in workbook.sheetnames:
    expected['地区下载索引'] = len(master['links'])
for name, count in expected.items():
    sheet=workbook[name]
    assert sheet['B5'].value == count, (name, sheet['B5'].value, count)
    if count:
        assert sum(sheet.cell(r,1).value is not None for r in range(9, sheet.max_row+1)) == count, name
    assert sheet.freeze_panes, name
assert [workbook['原始映射'].cell(r,1).value for r in range(9,90)] == [f'S{i:03}' for i in range(1,82)]
hyperlinks=[]
errors=[]
for sheet in workbook:
    for row in sheet:
        for cell in row:
            if cell.hyperlink:
                assert cell.hyperlink.target.startswith(('https://','http://'))
                hyperlinks.append((sheet.title,cell.coordinate,cell.hyperlink.target))
            if cell.data_type == 'e' or (isinstance(cell.value,str) and 'HYPERLINK is not implemented' in cell.value):
                errors.append((sheet.title,cell.coordinate,cell.value))
            assert not (isinstance(cell.value,str) and '[object Object]' in cell.value), (sheet.title,cell.coordinate)
assert not errors, errors
native=json.loads((root/'native-hyperlinks.json').read_text(encoding='utf-8'))
assert len(hyperlinks)==len(native)
assert master['coverage']['verified_complete_pdf_links']==sum(l.get('complete_pdf') is True for l in master['links'])
complete_links=[link for link in master['links'] if link.get('complete_pdf')]
assert master['coverage'].get('verified_complete_unique_sha256')==len({link['sha256'] for link in complete_links})
complete_sheet=workbook['已验完整PDF']
assert collections.Counter(complete_sheet.cell(r,3).hyperlink.target for r in range(9,9+len(complete_links)))==collections.Counter(link['final_url'] for link in complete_links)
assert collections.Counter(complete_sheet.cell(r,17).value for r in range(9,9+len(complete_links)))==collections.Counter(link['sha256'] for link in complete_links)
for link in master['links']:
    if link.get('complete_pdf'):
        assert not link.get('metadata_only')
        assert link.get('version_verified') and link.get('final_url') and link.get('sha256') and link.get('pages') and link.get('pdf_bytes') and link.get('audit_id') and link.get('audit_evidence_paths'), link['link_id']
draft_links=[link for link in master['links'] if link.get('draft_content_complete') and link.get('version_verified') and not link.get('complete_pdf')]
if draft_links:
    draft_sheet=workbook['完整草案PDF']
    assert collections.Counter(draft_sheet.cell(r,3).hyperlink.target for r in range(9,9+len(draft_links)))==collections.Counter(link['final_url'] for link in draft_links)
    assert not {link['link_id'] for link in draft_links} & {link['link_id'] for link in complete_links}
    for row, link in enumerate(draft_links,9):
        assert draft_sheet.cell(row,20).value == link['version_class']
        assert '内容完整' in draft_sheet.cell(row,21).value
        assert '不计正式' in draft_sheet.cell(row,22).value
    link_rows={workbook['链接明细'].cell(r,1).value:r for r in range(9,9+len(master['links']))}
    for link in draft_links:
        row=link_rows[link['link_id']]
        assert workbook['链接明细'].cell(row,19).value == '完整草案（非正式全文）'
standards={s['standard_id']:s for s in master['standards']}
if 'CECS-102-2002-FORMAL-E2012' in standards:
    edition=standards['CECS-102-2002-FORMAL-E2012']
    assert edition['year']==2002 and edition['edition_label']=='2012年版'
    assert not set(edition['link_ids']) & set(standards['CECS-102-2002-FORMAL']['link_ids'])
    assert edition['published_date'] is None and edition['effective_date'] is None
if 'CECS-796-2021-FORMAL' in standards:
    assert standards['CECS-796-2021-FORMAL']['name']=='屋面结构雪荷载设计标准'
if master['scope'].get('checkpoint') in ('1350','1410','1510'):
    assert standards['CECS-22-2005-FORMAL']['withdrawal_date']=='2025-04-01'
    assert standards['CECS-22-2005-FORMAL']['current_status'].startswith('已废止')
    standard_rows={workbook['规范总表'].cell(r,17).value:r for r in range(9,9+len(master['standards']))}
    row=standard_rows['CECS-22-2005-FORMAL']
    assert workbook['规范总表'].cell(row,5).value.startswith('已废止')
    assert workbook['规范总表'].cell(row,24).value.date().isoformat()=='2025-04-01'
    assert '第5页' in workbook['规范总表'].cell(row,25).value
if master['scope'].get('checkpoint') in ('1410','1510'):
    assert len(draft_links)==master['coverage']['verified_complete_draft_pdf_links']
    assert {l['version_class'] for l in draft_links}==({'报批稿','送审稿'} if master['scope']['checkpoint']=='1410' else {'报批稿','送审稿','征求意见稿'})
    for d in master['external']['ttbz_public_dom']['standards']:
        row=standard_rows[next(s['standard_id'] for s in master['standards'] if s.get('public_announcement_date')==d['public_announcement_date'] and s['year']==int(d['year']) and s['standard_number'].replace(' ','')==d['standard_number'].replace(' ',''))]
        assert workbook['规范总表'].cell(row,5).value=='现行（全国团体平台标注）'
        assert workbook['规范总表'].cell(row,26).value.date().isoformat()==d['public_announcement_date']
    for sid in ('CECS-1137-2022-FORMAL','CECS-22-2012-FORMAL'):
        row=standard_rows[sid]
        assert '待核' in workbook['规范总表'].cell(row,2).value
if master['scope'].get('checkpoint')=='1510':
    assert len(workbook.worksheets)==11
    index=workbook['地区下载索引']
    rows={index.cell(r,14).value:r for r in range(9,9+len(master['links']))}
    assert len(rows)==len(master['links'])
    region_ids={e['region_evidence_id'] for e in master['external']['download_source_region_registry']['evidence']}
    for link in master['links']:
        region=link['download_source_region']; row=rows[link['link_id']]; lr=link_rows[link['link_id']]
        assert index.cell(row,1).value==(region['province'] or '来源省份待核')
        assert index.cell(row,2).value==(region['city'] or '来源城市待核')
        assert index.cell(row,6).hyperlink.target==(link.get('final_url') or link['requested_url'])
        assert workbook['链接明细'].cell(lr,48).value==region['province']
        assert workbook['链接明细'].cell(lr,49).value==region['city']
        assert workbook['链接明细'].cell(lr,51).value==region['region_status']
        assert set(region['region_evidence_ids'])<=region_ids
        if region['province'] or region['city']:
            assert index.cell(row,12).value and index.cell(row,13).value
        assert workbook['链接明细'].cell(lr,44).value==link.get('body_reading_recommendation')
        assert workbook['链接明细'].cell(lr,45).value==link.get('body_searchable')
    for standard in master['standards']:
        assert workbook['规范总表'].cell(standard_rows[standard['standard_id']],27).value==standard['download_source_region_summary']
    old69=standards['CECS-69-1994-FORMAL']
    assert old69['withdrawal_date']=='2011-10-01' and old69['current_status'].startswith('已废止')
receipt={'checked_at_utc':__import__('datetime').datetime.now(__import__('datetime').timezone.utc).isoformat(),'xlsx_sha256':hashlib.sha256(book.read_bytes()).hexdigest(),'bytes':book.stat().st_size,'sheets':[{ 'name':s.title,'rows':s.max_row,'columns':s.max_column,'freeze_panes':s.freeze_panes} for s in workbook], 'native_hyperlinks':len(hyperlinks),'error_cells':len(errors),'all_81_original_mappings':True,'counts_reconciled':True,'complete_pdf_required_fields':True}
(root/'workbook-validation.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False))
